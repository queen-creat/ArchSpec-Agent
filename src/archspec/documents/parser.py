"""把不同格式的规范文件统一转换成可检索的纯文字。"""

from dataclasses import dataclass
from pathlib import Path

import pymupdf
from docx import Document
from docx.table import Table

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
SUPPORTED_MEDIA_TYPES = {"text/plain", DOCX_MEDIA_TYPE, "application/pdf"}


class DocumentParseError(ValueError):
    """所有可向用户说明的文档解析错误的基类。"""


class UnsupportedDocumentType(DocumentParseError):
    """上传格式不在第一阶段支持范围内。"""


class DocumentContainsNoText(DocumentParseError):
    """文件能够打开，但没有提取到可供检索的文字。"""


@dataclass(frozen=True)
class SourceBlock:
    """解析后的最小来源块；PDF 块带真实页码，其他格式页码未知。"""

    text: str
    page_number: int | None
    block_index: int


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    page_count: int | None
    blocks: tuple[SourceBlock, ...]


class DocumentParser:
    """Parser（解析器）负责把 TXT、DOCX、PDF 翻译成同一种文字格式。"""

    def parse(self, path: Path, media_type: str) -> ParsedDocument:
        if media_type == "text/plain":
            text = path.read_text(encoding="utf-8-sig").strip()
            blocks = (SourceBlock(text, None, 0),) if text else ()
            result = ParsedDocument(text, None, blocks)
        elif media_type == DOCX_MEDIA_TYPE:
            blocks = self._read_docx(path)
            result = ParsedDocument(
                "\n".join(block.text for block in blocks),
                None,
                blocks,
            )
        elif media_type == "application/pdf":
            try:
                result = self._read_pdf(path)
            except (pymupdf.FileDataError, RuntimeError, ValueError):
                raise DocumentParseError("PDF无法提取文字，请对照原件查看") from None
        else:
            raise UnsupportedDocumentType(f"不支持的文件类型：{media_type}")

        if not result.text.strip():
            raise DocumentContainsNoText("文件中没有提取到文字，可能是扫描版 PDF 或空文件")
        return result

    @staticmethod
    def _read_docx(path: Path) -> tuple[SourceBlock, ...]:
        texts: list[str] = []
        document = Document(path)
        # iter_inner_content 按 Word 中真实出现的顺序遍历段落和表格。
        for item in document.iter_inner_content():
            if isinstance(item, Table):
                for row in item.rows:
                    for cell in row.cells:
                        if text := cell.text.strip():
                            texts.append(text)
            elif text := item.text.strip():
                texts.append(text)
        return tuple(
            SourceBlock(text=text, page_number=None, block_index=index)
            for index, text in enumerate(texts)
        )

    @staticmethod
    def _read_pdf(path: Path) -> ParsedDocument:
        with pymupdf.open(path) as document:
            blocks: list[SourceBlock] = []
            for page_number, page in enumerate(document, start=1):
                if text := page.get_text().strip():
                    blocks.append(
                        SourceBlock(
                            text=text,
                            page_number=page_number,
                            block_index=len(blocks),
                        )
                    )
            return ParsedDocument(
                text="\n".join(block.text for block in blocks),
                page_count=document.page_count,
                blocks=tuple(blocks),
            )
